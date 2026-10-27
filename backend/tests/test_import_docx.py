from docx import Document

from scripts.import_docx import parse_docx, save


def make_docx(path):
    """Qog'oz ma'lumotnomaga o'xshash jadval: ikki qatorli sarlavha, birlashgan bo'lim qatorlari."""
    doc = Document()
    doc.add_paragraph("TELEFON MA'LUMOTNOMASI")
    table = doc.add_table(rows=0, cols=6)

    def row(*values):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
        return cells

    head = row("№", "Ф.И.Ш.", "Лавозими", "Хона", "Телефон", "")
    head[4].merge(head[5])
    head[4].text = "Телефон"
    row("", "", "", "", "Ички", "Мобил")
    dept = row("Rahbariyat", "", "", "", "", "")
    dept[0].merge(dept[5]).text = "Rahbariyat"
    row("1", "Karimov Anvar\nAliyevich", "Vazir", "201", "100", "90 123 45 67")
    row("2", "Toshmatova Gulnora", "Vazir o'rinbosari", "203", "101\n102", "")
    dept = row("", "", "", "", "", "")
    dept[0].merge(dept[5]).text = "Moliya bo'limi"
    row("3", "Rahimov Bobur", "Bosh hisobchi", "310", "150", "")
    # Nomida ustun kalit so'zi bor bo'limlar ("ichki", "xona") sarlavha deb olinmasligi kerak.
    for name, person in (("Ichki audit bo'limi", "Aliyev Ali"), ("Ijro intizomi va devonxona bo'limi", "Valiyeva Vali")):
        dept = row("", "", "", "", "", "")
        dept[0].merge(dept[5]).text = name
        row("4", person, "Bosh mutaxassis", "401", "160", "")
    doc.save(path)


def test_parse_docx(tmp_path):
    path = tmp_path / "eski.docx"
    make_docx(path)
    columns, sections = parse_docx(path)
    assert set(columns.values()) == {"full_name", "position", "room", "internal_phone", "mobile_phone"}
    assert [(name, len(emps)) for name, emps in sections] == [
        ("Rahbariyat", 2),
        ("Moliya bo'limi", 1),
        ("Ichki audit bo'limi", 1),
        ("Ijro intizomi va devonxona bo'limi", 1),
    ]
    first, second = sections[0][1]
    assert first["full_name"] == "Karimov Anvar Aliyevich"
    assert first["mobile_phone"] == "90 123 45 67"
    assert second["internal_phone"] == "101, 102"
    audit = sections[2][1][0]
    assert (audit["full_name"], audit["room"], audit["internal_phone"]) == ("Aliyev Ali", "401", "160")
    devon = sections[3][1][0]
    assert (devon["full_name"], devon["position"], devon["internal_phone"]) == ("Valiyeva Vali", "Bosh mutaxassis", "160")


def test_save_refuses_to_overwrite(tmp_path, db):
    path = tmp_path / "eski.docx"
    make_docx(path)
    _, sections = parse_docx(path)
    save(db, sections, replace=False)
    assert db.execute("SELECT COUNT(*) FROM employees").fetchone()[0] == 5
    try:
        save(db, sections, replace=False)
    except SystemExit:
        pass
    else:
        raise AssertionError("ikkinchi import rad etilishi kerak edi")
    save(db, sections, replace=True)
    assert db.execute("SELECT COUNT(*) FROM employees").fetchone()[0] == 5
