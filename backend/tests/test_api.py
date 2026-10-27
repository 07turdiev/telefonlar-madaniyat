from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import add_user, login


def create_department(client, name, parent_id=None):
    response = client.post("/api/admin/departments", json={"name": name, "parent_id": parent_id})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def create_employee(client, department_id, full_name, **fields):
    response = client.post(
        "/api/admin/employees",
        json={"department_id": department_id, "full_name": full_name, **fields},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def section_names(client, url="/api/directory"):
    return [(s["name"], [e["full_name"] for e in s["employees"]]) for s in client.get(url).json()["sections"]]


# --- Ochiq sahifa ----------------------------------------------------------


def test_directory_is_public_and_hides_empty_departments(hr):
    leaders = create_department(hr, "Rahbariyat")
    culture = create_department(hr, "Madaniyat boshqarmasi")
    create_department(hr, "Bo'sh bo'lim")
    theatre = create_department(hr, "Teatr bo'limi", parent_id=culture)
    create_employee(hr, leaders, "Karimov Anvar", position="Vazir", internal_phone="1000")
    create_employee(hr, theatre, "Aliyeva Dilnoza", internal_phone="1215", room="305")

    hr.post("/api/auth/logout")
    assert hr.get("/api/auth/me").status_code == 401

    data = hr.get("/api/directory").json()
    assert data["org_name"] == "Test vazirligi"
    assert data["updated"]
    assert [(s["name"], s["level"]) for s in data["sections"]] == [
        ("Rahbariyat", 0),
        ("Madaniyat boshqarmasi", 0),  # o'zida hodim yo'q, lekin ichki bo'limida bor
        ("Teatr bo'limi", 1),
    ]
    assert data["sections"][2]["employees"][0]["room"] == "305"


def test_updated_ignores_account_actions(hr, db):
    create_department(hr, "Moliya")
    create_employee(hr, db.execute("SELECT id FROM departments").fetchone()[0], "Ali")
    before = hr.get("/api/directory").json()["updated"]
    db.execute(
        "INSERT INTO audit_log (at, username, action, details)"
        " VALUES ('2999-01-01T00:00:00', 'admin', 'parol almashtirildi', 'kadrlar')"
    )
    db.commit()
    assert hr.get("/api/directory").json()["updated"] == before


def test_admin_directory_includes_empty_departments(hr):
    create_department(hr, "Bo'sh bo'lim")
    assert section_names(hr, "/api/admin/directory") == [("Bo'sh bo'lim", [])]


# --- Kirish va himoya ------------------------------------------------------


def test_admin_requires_login(client):
    assert client.get("/api/admin/directory").status_code == 401
    assert client.post("/api/admin/departments", json={"name": "X"}).status_code == 401


def test_wrong_password_and_lockout(client, db):
    add_user(db, "kadrlar")
    for _ in range(5):
        r = client.post("/api/auth/login", json={"username": "kadrlar", "password": "xato"})
        assert r.status_code == 400
        assert r.json()["detail"] == "Login yoki parol noto'g'ri."
    r = client.post("/api/auth/login", json={"username": "kadrlar", "password": "kadrlar-parol"})
    assert r.status_code == 429


def test_me_does_not_leak_password_hash(hr):
    me = hr.get("/api/auth/me").json()
    assert me["username"] == "kadrlar"
    assert "password_hash" not in me


def test_post_without_xhr_header_is_rejected(hr):
    del hr.headers["X-Requested-With"]
    r = hr.post("/api/admin/departments", json={"name": "Boshqa saytdan"})
    assert r.status_code == 403
    hr.headers["X-Requested-With"] = "XMLHttpRequest"
    assert hr.get("/api/admin/departments").json() == []


def test_change_own_password(hr):
    r = hr.post("/api/auth/password", json={"current": "xato", "password": "yangi-parol-123"})
    assert r.json()["detail"]["errors"]["current"]
    r = hr.post("/api/auth/password", json={"current": "kadrlar-parol", "password": "qisqa"})
    assert r.json()["detail"]["errors"]["password"]
    r = hr.post("/api/auth/password", json={"current": "kadrlar-parol", "password": "yangi-parol-123"})
    assert r.status_code == 200
    hr.post("/api/auth/logout")
    r = hr.post("/api/auth/login", json={"username": "kadrlar", "password": "yangi-parol-123"})
    assert r.status_code == 200


# --- Hodimlar --------------------------------------------------------------


def test_employee_validation(hr):
    dept = create_department(hr, "Moliya")
    r = hr.post("/api/admin/employees", json={"department_id": dept, "full_name": "  ", "email": "xato"})
    assert r.status_code == 400
    assert set(r.json()["detail"]["errors"]) == {"full_name", "email"}
    r = hr.post("/api/admin/employees", json={"department_id": 999, "full_name": "Ali"})
    assert set(r.json()["detail"]["errors"]) == {"department_id"}


def test_internal_phone_is_single_four_digit_number(hr):
    dept = create_department(hr, "Moliya")
    for bad in ("1000, 68-11", "68-11", "123", "12345", "12a4"):
        r = hr.post("/api/admin/employees", json={"department_id": dept, "full_name": "Ali", "internal_phone": bad})
        assert set(r.json()["detail"]["errors"]) == {"internal_phone"}, bad
    create_employee(hr, dept, "Ali", internal_phone=" 1290 ")
    create_employee(hr, dept, "Vali", internal_phone="")  # bo'sh qoldirish mumkin
    assert [e["internal_phone"] for e in hr.get("/api/directory").json()["sections"][0]["employees"]] == ["1290", ""]


def test_employee_update_logs_changes(hr):
    finance = create_department(hr, "Moliya")
    legal = create_department(hr, "Yuridik")
    emp = create_employee(hr, finance, "Karimov  Anvar", internal_phone="1230")
    r = hr.put(
        f"/api/admin/employees/{emp}",
        json={"department_id": legal, "full_name": "Karimov Anvar", "internal_phone": "1250"},
    )
    assert r.status_code == 200
    assert r.json()["full_name"] == "Karimov Anvar"
    entry = hr.get("/api/admin/log").json()[0]
    assert entry["username"] == "kadrlar"
    assert "«1230» → «1250»" in entry["details"]
    assert "«Moliya» → «Yuridik»" in entry["details"]


def test_employee_order_and_delete(hr):
    dept = create_department(hr, "Moliya")
    a = create_employee(hr, dept, "A")
    create_employee(hr, dept, "B")
    c = create_employee(hr, dept, "C")
    hr.post(f"/api/admin/employees/{c}/move", json={"direction": "up"})
    assert section_names(hr) == [("Moliya", ["A", "C", "B"])]
    hr.post(f"/api/admin/employees/{a}/move", json={"direction": "up"})  # birinchisi joyida qoladi
    assert section_names(hr) == [("Moliya", ["A", "C", "B"])]
    assert hr.delete(f"/api/admin/employees/{a}").status_code == 200
    assert section_names(hr) == [("Moliya", ["C", "B"])]
    assert hr.delete(f"/api/admin/employees/{a}").status_code == 404


# --- Bo'limlar -------------------------------------------------------------


def test_department_nesting_rules(hr):
    top = create_department(hr, "Boshqarma")
    child = create_department(hr, "Bo'lim", parent_id=top)
    r = hr.post("/api/admin/departments", json={"name": "Sektor", "parent_id": child})
    assert r.json()["detail"]["errors"]["parent_id"]  # faqat ikki daraja
    other = create_department(hr, "Boshqa")
    r = hr.put(f"/api/admin/departments/{top}", json={"name": "Boshqarma", "parent_id": other})
    assert r.json()["detail"]["errors"]["parent_id"]  # ichida bo'limi bor


def test_department_delete_refused_when_not_empty(hr):
    dept = create_department(hr, "Moliya")
    emp = create_employee(hr, dept, "Ali")
    r = hr.delete(f"/api/admin/departments/{dept}")
    assert r.status_code == 400
    assert "hodimlar bor" in r.json()["detail"]
    hr.delete(f"/api/admin/employees/{emp}")
    assert hr.delete(f"/api/admin/departments/{dept}").status_code == 200


def test_department_counts_and_move(hr):
    create_department(hr, "A")
    b = create_department(hr, "B")
    create_employee(hr, b, "Ali")
    hr.post(f"/api/admin/departments/{b}/move", json={"direction": "up"})
    assert [(d["name"], d["count"]) for d in hr.get("/api/admin/departments").json()] == [
        ("B", 1),
        ("A", 0),
    ]


# --- Foydalanuvchilar ------------------------------------------------------


def test_only_admin_manages_users(hr):
    assert hr.get("/api/admin/users").status_code == 403


def test_admin_creates_and_deletes_users(admin):
    r = admin.post("/api/admin/users", json={"username": "a b", "password": "123"})
    assert set(r.json()["detail"]["errors"]) == {"username", "password"}
    r = admin.post(
        "/api/admin/users",
        json={"username": "kadrlar2", "full_name": "Kadrlar", "password": "uzun-parol"},
    )
    assert r.status_code == 201
    new_id = r.json()["id"]
    r = admin.post("/api/admin/users", json={"username": "KADRLAR2", "password": "uzun-parol"})
    assert r.json()["detail"]["errors"]["username"] == "Bunday login bor."
    me = admin.get("/api/auth/me").json()
    assert admin.delete(f"/api/admin/users/{me['id']}").status_code == 400
    assert admin.delete(f"/api/admin/users/{new_id}").status_code == 200


def test_admin_resets_password(admin, db):
    add_user(db, "kadrlar")
    target = db.execute("SELECT id FROM users WHERE username = 'kadrlar'").fetchone()["id"]
    r = admin.post(f"/api/admin/users/{target}/password", json={"password": "yangi-parol-1"})
    assert r.status_code == 200
    admin.post("/api/auth/logout")
    r = admin.post("/api/auth/login", json={"username": "kadrlar", "password": "yangi-parol-1"})
    assert r.status_code == 200


# --- Eksport ---------------------------------------------------------------


def test_export_xlsx(hr):
    dept = create_department(hr, "Moliya")
    create_employee(hr, dept, "Karimov Anvar", position="Bosh hisobchi", internal_phone="1230")
    r = hr.get("/api/admin/export.xlsx")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    assert [c.value for c in ws[2]][:5] == ["Moliya", "Karimov Anvar", "Bosh hisobchi", None, "1230"]
