"""Kadrlar bo'limi API'si: /api/admin/*. Hamma endpoint login talab qiladi."""

import io
import re
import sqlite3
from datetime import date
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel

from app.auth import MIN_PASSWORD_LENGTH, hash_password
from app.db import (
    EMPLOYEE_FIELDS,
    VACANT_CLEARED,
    VACANT_NAME,
    department_tree,
    directory,
    log,
    move,
    next_sort_order,
    now,
)
from app.deps import admin_user, current_user, get_db, user_out

router = APIRouter(prefix="/api/admin", dependencies=[Depends(current_user)])

FIELD_LIMIT = 200
LABELS = {
    "full_name": "F.I.Sh.",
    "position": "lavozim",
    "room": "xona",
    "internal_phone": "ichki tel.",
    "city_phone": "shahar tel.",
    "mobile_phone": "mobil",
    "email": "e-mail",
    "department_id": "bo'lim",
    "is_vacant": "vakant",
}
EMPLOYEE_COLUMNS = [*EMPLOYEE_FIELDS, "department_id", "is_vacant"]
USERNAME_RE = re.compile(r"[A-Za-z0-9._-]{3,32}")
EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")
INTERNAL_PHONE_RE = re.compile(r"\d{4}")
PASSWORD_ERROR = f"Parol kamida {MIN_PASSWORD_LENGTH} belgidan iborat bo'lsin."


class EmployeeIn(BaseModel):
    department_id: int | None = None
    full_name: str = ""
    position: str = ""
    room: str = ""
    internal_phone: str = ""
    city_phone: str = ""
    mobile_phone: str = ""
    email: str = ""
    is_vacant: bool = False


class DepartmentIn(BaseModel):
    name: str = ""
    parent_id: int | None = None


class MoveIn(BaseModel):
    direction: Literal["up", "down"]


class UserIn(BaseModel):
    username: str = ""
    full_name: str = ""
    password: str = ""
    is_admin: bool = False


class NewPasswordIn(BaseModel):
    password: str


def clean(value: str) -> str:
    return " ".join(value.split())[:FIELD_LIMIT]


def field_errors(errors: dict[str, str]) -> None:
    if errors:
        raise HTTPException(400, detail={"errors": errors})


def get_or_404(db, table: str, item_id: int) -> sqlite3.Row:
    row = db.execute(f"SELECT * FROM {table} WHERE id = ?", (item_id,)).fetchone()
    if row is None:
        raise HTTPException(404, "Topilmadi.")
    return row


def department_names(db) -> dict[int, str]:
    return {r["id"]: r["name"] for r in db.execute("SELECT id, name FROM departments")}


# --- Hodimlar --------------------------------------------------------------


def _validated_employee(data: EmployeeIn, db) -> dict:
    values = {field: clean(getattr(data, field)) for field in EMPLOYEE_FIELDS}
    values["department_id"] = data.department_id
    values["is_vacant"] = int(data.is_vacant)
    errors = {}
    if data.is_vacant:
        values["full_name"] = VACANT_NAME
        values.update({field: "" for field in VACANT_CLEARED})
        if not values["position"]:
            errors["position"] = "Vakant o'rin uchun lavozimni kiriting."
    if not values["full_name"]:
        errors["full_name"] = "F.I.Sh. kiritilishi shart."
    if values["internal_phone"] and not INTERNAL_PHONE_RE.fullmatch(values["internal_phone"]):
        errors["internal_phone"] = "Ichki raqam bitta 4 xonali raqam bo'lsin, masalan 1290."
    if values["email"] and not EMAIL_RE.fullmatch(values["email"]):
        errors["email"] = "E-mail manzili noto'g'ri yozilgan."
    if data.department_id is None or not db.execute(
        "SELECT 1 FROM departments WHERE id = ?", (data.department_id,)
    ).fetchone():
        errors["department_id"] = "Bo'limni tanlang."
    field_errors(errors)
    return values


@router.get("/directory")
def admin_directory(db=Depends(get_db)):
    """Bo'sh bo'limlar ham kiradi: kadrlar bo'limi ularga hodim qo'shadi."""
    return {"sections": directory(db, include_empty=True)}


@router.post("/employees", status_code=201)
def employee_create(data: EmployeeIn, db=Depends(get_db), user=Depends(current_user)):
    values = _validated_employee(data, db)
    cursor = db.execute(
        f"INSERT INTO employees ({', '.join(EMPLOYEE_COLUMNS)}, sort_order, updated_at)"
        f" VALUES ({', '.join('?' * len(EMPLOYEE_COLUMNS))}, ?, ?)",
        [values[c] for c in EMPLOYEE_COLUMNS]
        + [next_sort_order(db, "employees", "department_id", values["department_id"]), now()],
    )
    dept = department_names(db)[values["department_id"]]
    log(db, user["username"], "hodim qo'shildi", f"{values['full_name']} ({dept})")
    db.commit()
    return dict(get_or_404(db, "employees", cursor.lastrowid))


@router.put("/employees/{employee_id}")
def employee_update(
    employee_id: int, data: EmployeeIn, db=Depends(get_db), user=Depends(current_user)
):
    employee = get_or_404(db, "employees", employee_id)
    values = _validated_employee(data, db)

    names = department_names(db)
    changes = []
    for field, label in LABELS.items():
        old, new = employee[field], values[field]
        if old != new:
            if field == "department_id":
                old, new = names.get(old, "?"), names[new]
            elif field == "is_vacant":
                old, new = ("ha" if old else "yo'q"), ("ha" if new else "yo'q")
            changes.append(f"{label}: «{old}» → «{new}»")
    if not changes:
        return dict(employee)

    sort_order = employee["sort_order"]
    if values["department_id"] != employee["department_id"]:
        sort_order = next_sort_order(db, "employees", "department_id", values["department_id"])
    db.execute(
        f"UPDATE employees SET {', '.join(f'{c} = ?' for c in EMPLOYEE_COLUMNS)},"
        " sort_order = ?, updated_at = ? WHERE id = ?",
        [values[c] for c in EMPLOYEE_COLUMNS] + [sort_order, now(), employee_id],
    )
    log(db, user["username"], "hodim o'zgartirildi", f"{employee['full_name']}: " + "; ".join(changes))
    db.commit()
    return dict(get_or_404(db, "employees", employee_id))


@router.delete("/employees/{employee_id}")
def employee_delete(employee_id: int, db=Depends(get_db), user=Depends(current_user)):
    employee = get_or_404(db, "employees", employee_id)
    dept = department_names(db).get(employee["department_id"], "?")
    db.execute("DELETE FROM employees WHERE id = ?", (employee_id,))
    log(db, user["username"], "hodim o'chirildi", f"{employee['full_name']} ({dept})")
    db.commit()
    return {"ok": True}


@router.post("/employees/{employee_id}/move")
def employee_move(employee_id: int, data: MoveIn, db=Depends(get_db), user=Depends(current_user)):
    employee = get_or_404(db, "employees", employee_id)
    if move(db, "employees", "department_id", employee_id, data.direction):
        log(db, user["username"], "tartib o'zgartirildi", employee["full_name"])
        db.commit()
    return {"ok": True}


# --- Bo'limlar -------------------------------------------------------------


def _validated_department(data: DepartmentIn, db, department=None) -> dict:
    values = {"name": clean(data.name), "parent_id": data.parent_id}
    errors = {}
    if not values["name"]:
        errors["name"] = "Bo'lim nomini kiriting."
    if data.parent_id is not None:
        parent = db.execute("SELECT * FROM departments WHERE id = ?", (data.parent_id,)).fetchone()
        if parent is None or parent["parent_id"] is not None:
            errors["parent_id"] = "Faqat yuqori darajadagi bo'lim tanlanadi."
        elif department is not None:
            if data.parent_id == department["id"]:
                errors["parent_id"] = "Bo'lim o'zining ichida bo'la olmaydi."
            elif db.execute(
                "SELECT 1 FROM departments WHERE parent_id = ?", (department["id"],)
            ).fetchone():
                errors["parent_id"] = (
                    "Bu bo'lim ichida boshqa bo'limlar bor, uni boshqa bo'lim ichiga qo'yib bo'lmaydi."
                )
    field_errors(errors)
    return values


@router.get("/departments")
def departments(db=Depends(get_db)):
    counts = dict(
        db.execute("SELECT department_id, COUNT(*) FROM employees GROUP BY department_id").fetchall()
    )
    return [{**d, "count": counts.get(d["id"], 0)} for d in department_tree(db)]


@router.post("/departments", status_code=201)
def department_create(data: DepartmentIn, db=Depends(get_db), user=Depends(current_user)):
    values = _validated_department(data, db)
    cursor = db.execute(
        "INSERT INTO departments (name, parent_id, sort_order) VALUES (?, ?, ?)",
        (
            values["name"],
            values["parent_id"],
            next_sort_order(db, "departments", "parent_id", values["parent_id"]),
        ),
    )
    log(db, user["username"], "bo'lim qo'shildi", values["name"])
    db.commit()
    return dict(get_or_404(db, "departments", cursor.lastrowid))


@router.put("/departments/{department_id}")
def department_update(
    department_id: int, data: DepartmentIn, db=Depends(get_db), user=Depends(current_user)
):
    department = get_or_404(db, "departments", department_id)
    values = _validated_department(data, db, department)
    sort_order = department["sort_order"]
    if values["parent_id"] != department["parent_id"]:
        sort_order = next_sort_order(db, "departments", "parent_id", values["parent_id"])
    db.execute(
        "UPDATE departments SET name = ?, parent_id = ?, sort_order = ? WHERE id = ?",
        (values["name"], values["parent_id"], sort_order, department_id),
    )
    details = values["name"]
    if values["name"] != department["name"]:
        details = f"«{department['name']}» → «{values['name']}»"
    log(db, user["username"], "bo'lim o'zgartirildi", details)
    db.commit()
    return dict(get_or_404(db, "departments", department_id))


@router.delete("/departments/{department_id}")
def department_delete(department_id: int, db=Depends(get_db), user=Depends(current_user)):
    department = get_or_404(db, "departments", department_id)
    name = department["name"]
    if db.execute("SELECT 1 FROM employees WHERE department_id = ?", (department_id,)).fetchone():
        raise HTTPException(
            400, f"«{name}» ichida hodimlar bor. Avval ularni o'chiring yoki boshqa bo'limga o'tkazing."
        )
    if db.execute("SELECT 1 FROM departments WHERE parent_id = ?", (department_id,)).fetchone():
        raise HTTPException(
            400, f"«{name}» ichida boshqa bo'limlar bor. Avval ularni o'chiring yoki ko'chiring."
        )
    db.execute("DELETE FROM departments WHERE id = ?", (department_id,))
    log(db, user["username"], "bo'lim o'chirildi", name)
    db.commit()
    return {"ok": True}


@router.post("/departments/{department_id}/move")
def department_move(
    department_id: int, data: MoveIn, db=Depends(get_db), user=Depends(current_user)
):
    department = get_or_404(db, "departments", department_id)
    if move(db, "departments", "parent_id", department_id, data.direction):
        log(db, user["username"], "tartib o'zgartirildi", department["name"])
        db.commit()
    return {"ok": True}


# --- Foydalanuvchilar (faqat administrator) --------------------------------


@router.get("/users")
def users(db=Depends(get_db), user=Depends(admin_user)):
    return [user_out(u) for u in db.execute("SELECT * FROM users ORDER BY username")]


@router.post("/users", status_code=201)
def user_create(data: UserIn, db=Depends(get_db), user=Depends(admin_user)):
    username = data.username.strip()
    errors = {}
    if not USERNAME_RE.fullmatch(username):
        errors["username"] = "Login 3–32 belgi: lotin harflari, raqamlar, nuqta, tire."
    elif db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
        errors["username"] = "Bunday login bor."
    if len(data.password) < MIN_PASSWORD_LENGTH:
        errors["password"] = PASSWORD_ERROR
    field_errors(errors)
    cursor = db.execute(
        "INSERT INTO users (username, full_name, password_hash, is_admin, created_at)"
        " VALUES (?, ?, ?, ?, ?)",
        (username, clean(data.full_name), hash_password(data.password), int(data.is_admin), now()),
    )
    log(db, user["username"], "foydalanuvchi qo'shildi", username)
    db.commit()
    return user_out(get_or_404(db, "users", cursor.lastrowid))


@router.post("/users/{user_id}/password")
def user_reset_password(
    user_id: int, data: NewPasswordIn, db=Depends(get_db), user=Depends(admin_user)
):
    target = get_or_404(db, "users", user_id)
    if len(data.password) < MIN_PASSWORD_LENGTH:
        field_errors({"password": PASSWORD_ERROR})
    db.execute(
        "UPDATE users SET password_hash = ? WHERE id = ?", (hash_password(data.password), user_id)
    )
    log(db, user["username"], "parol almashtirildi", target["username"])
    db.commit()
    return {"ok": True}


@router.delete("/users/{user_id}")
def user_delete(user_id: int, db=Depends(get_db), user=Depends(admin_user)):
    target = get_or_404(db, "users", user_id)
    if target["id"] == user["id"]:
        raise HTTPException(400, "O'zingizni o'chira olmaysiz.")
    db.execute("DELETE FROM users WHERE id = ?", (user_id,))
    log(db, user["username"], "foydalanuvchi o'chirildi", target["username"])
    db.commit()
    return {"ok": True}


# --- Jurnal va eksport -----------------------------------------------------


@router.get("/log")
def audit_log(db=Depends(get_db)):
    return [dict(r) for r in db.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 500")]


@router.get("/export.xlsx")
def export_xlsx(db=Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Telefonlar"
    ws.append(["Bo'lim", "F.I.Sh.", "Lavozimi", "Xona", "Ichki", "Shahar tel.", "Mobil", "E-mail"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE6F0")
    for section in directory(db):
        for emp in section["employees"]:
            vacant = emp["is_vacant"]
            ws.append(
                [section["name"], *("-" if vacant and f in VACANT_CLEARED else emp[f] for f in EMPLOYEE_FIELDS)]
            )
    for column, width in zip("ABCDEFGH", (34, 34, 40, 8, 10, 16, 18, 28)):
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="telefonlar-{date.today():%Y-%m-%d}.xlsx"'
        },
    )
